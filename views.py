import io
import json

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, logout
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse, Http404
from django.views.decorators.http import require_POST
from django.db.models import Count

from .forms import RegisterForm, LoginForm
from .models import (
    AnalysisSession, RequirementGap, TestScenario,
    FeedbackItem, TestCondition, DetailedTestCase,
    Workspace, WorkspaceMembership,
)
from .services.qa_analyzer import QAAnalyzer
from .services.detailed_case_generator import DetailedCaseGenerator
from .services.training_data_manager import TrainingDataManager
from .services.gemini_service import GeminiQuotaExceeded


# ─────────────────────────────────────────────
# Public pages
# ─────────────────────────────────────────────

def home_view(request):
    return render(request, 'home.html')


def about_view(request):
    return render(request, 'about.html')


# ─────────────────────────────────────────────
# Authentication
# ─────────────────────────────────────────────

def register_view(request):
    if request.user.is_authenticated:
        return redirect('workspace')

    if request.method == 'POST':
        form = RegisterForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user, backend='django.contrib.auth.backends.ModelBackend')
            return redirect('workspace')
    else:
        form = RegisterForm()

    return render(request, 'register.html', {'form': form})


def login_view(request):
    if request.user.is_authenticated:
        return redirect('workspace')

    if request.method == 'POST':
        form = LoginForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            login(request, user)
            return redirect('workspace')
    else:
        form = LoginForm()

    return render(request, 'login.html', {'form': form})


def logout_view(request):
    logout(request)
    return redirect('home')


# ─────────────────────────────────────────────
# Main workspace
# ─────────────────────────────────────────────

def _my_workspaces(user):
    """Team workspaces this user belongs to, newest first, with member counts."""
    ws_ids = WorkspaceMembership.objects.filter(user=user).values_list('workspace_id', flat=True)
    return Workspace.objects.filter(id__in=ws_ids).annotate(
        member_count=Count('memberships', distinct=True)
    ).order_by('-created_at')


@login_required
def workspace_view(request):
    history = AnalysisSession.objects.filter(
        user=request.user, workspace__isnull=True
    ).order_by('-is_pinned', '-created_at')
    return render(request, 'workspace.html', {
        'history': history,
        'my_workspaces': _my_workspaces(request.user),
        'can_delete_chats': True,  # personal chats are always the user's own
    })


# ─────────────────────────────────────────────
# Shared workspace (multi-user)
# ─────────────────────────────────────────────

@login_required
def shared_workspace_view(request, workspace_id):
    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)

    membership, just_joined = WorkspaceMembership.objects.get_or_create(
        workspace=workspace,
        user=request.user,
        defaults={'role': 'member'},
    )

    history = AnalysisSession.objects.filter(
        workspace=workspace
    ).order_by('-is_pinned', '-created_at')

    members = WorkspaceMembership.objects.filter(
        workspace=workspace
    ).select_related('user')

    return render(request, 'workspace.html', {
        'history': history,
        'workspace': workspace,
        'members': members,
        'just_joined': just_joined,
        'my_workspaces': _my_workspaces(request.user),
        'can_delete_chats': workspace.owner == request.user,
        'is_owner': workspace.owner == request.user,
    })


# ─────────────────────────────────────────────
# Create workspace
# ─────────────────────────────────────────────

@login_required
@require_POST
def create_workspace_view(request):
    try:
        body = json.loads(request.body)
        name = body.get('name', '').strip() or 'Team Workspace'
    except (json.JSONDecodeError, AttributeError):
        name = 'Team Workspace'

    workspace = Workspace.objects.create(name=name[:200], owner=request.user)
    WorkspaceMembership.objects.create(workspace=workspace, user=request.user, role='owner')

    return JsonResponse({'workspace_id': workspace.workspace_id, 'name': workspace.name})


# ─────────────────────────────────────────────
# Rename a workspace (owner only)
# ─────────────────────────────────────────────

@login_required
@require_POST
def rename_workspace_view(request):
    try:
        body = json.loads(request.body)
        workspace_id = body.get('workspace_id', '').strip()
        name = body.get('name', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not workspace_id or not name:
        return JsonResponse({'error': 'workspace_id and name are required.'}, status=400)

    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
    if workspace.owner != request.user:
        return JsonResponse({'error': 'Only the workspace owner can rename it.'}, status=403)

    workspace.name = name[:200]
    workspace.save(update_fields=['name'])
    return JsonResponse({'status': 'ok', 'name': workspace.name})


# ─────────────────────────────────────────────
# Leave / quit a workspace (any member)
# ─────────────────────────────────────────────

@login_required
@require_POST
def leave_workspace_view(request):
    try:
        body = json.loads(request.body)
        workspace_id = body.get('workspace_id', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not workspace_id:
        return JsonResponse({'error': 'workspace_id is required.'}, status=400)

    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
    WorkspaceMembership.objects.filter(workspace=workspace, user=request.user).delete()
    return JsonResponse({'status': 'left'})


# ─────────────────────────────────────────────
# Add a member by username (owner only)
# ─────────────────────────────────────────────

@login_required
@require_POST
def add_member_view(request):
    try:
        body = json.loads(request.body)
        workspace_id = body.get('workspace_id', '').strip()
        username = body.get('username', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not workspace_id or not username:
        return JsonResponse({'error': 'workspace_id and username are required.'}, status=400)

    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
    if workspace.owner != request.user:
        return JsonResponse({'error': 'Only the workspace owner can add members.'}, status=403)

    try:
        user = User.objects.get(username__iexact=username)
    except User.DoesNotExist:
        return JsonResponse({'error': 'No user named "%s".' % username}, status=404)

    _, created = WorkspaceMembership.objects.get_or_create(
        workspace=workspace, user=user, defaults={'role': 'member'}
    )
    if not created:
        return JsonResponse({'error': '%s is already a member.' % user.username}, status=400)

    return JsonResponse({
        'status': 'added',
        'username': user.username,
        'member_count': WorkspaceMembership.objects.filter(workspace=workspace).count(),
    })


# ─────────────────────────────────────────────
# Remove a member (owner only)
# ─────────────────────────────────────────────

@login_required
@require_POST
def remove_member_view(request):
    try:
        body = json.loads(request.body)
        workspace_id = body.get('workspace_id', '').strip()
        username = body.get('username', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not workspace_id or not username:
        return JsonResponse({'error': 'workspace_id and username are required.'}, status=400)

    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
    if workspace.owner != request.user:
        return JsonResponse({'error': 'Only the workspace owner can remove members.'}, status=403)
    if username.lower() == workspace.owner.username.lower():
        return JsonResponse({'error': 'The owner cannot be removed.'}, status=400)

    deleted, _ = WorkspaceMembership.objects.filter(
        workspace=workspace, user__username__iexact=username
    ).delete()
    if not deleted:
        return JsonResponse({'error': '%s is not a member.' % username}, status=404)

    return JsonResponse({
        'status': 'removed',
        'username': username,
        'member_count': WorkspaceMembership.objects.filter(workspace=workspace).count(),
    })


# ─────────────────────────────────────────────
# Delete a whole workspace + all its chats (owner only)
# ─────────────────────────────────────────────

@login_required
@require_POST
def delete_workspace_view(request):
    try:
        body = json.loads(request.body)
        workspace_id = body.get('workspace_id', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not workspace_id:
        return JsonResponse({'error': 'workspace_id is required.'}, status=400)

    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
    if workspace.owner != request.user:
        return JsonResponse({'error': 'Only the workspace owner can delete the workspace.'}, status=403)

    # Sessions use SET_NULL, so remove them explicitly for a complete deletion;
    # this cascades to their scenarios/conditions/gaps/feedback. Memberships
    # cascade automatically when the workspace is deleted.
    AnalysisSession.objects.filter(workspace=workspace).delete()
    workspace.delete()

    return JsonResponse({'status': 'deleted'})


# ─────────────────────────────────────────────
# Serialization helpers
# ─────────────────────────────────────────────

def _serialize_scenarios(session):
    return [
        {
            'db_id': s.id,
            'id': s.scenario_id,
            'requirement_ref': s.requirement_ref,
            'condition_ref': s.condition_ref,
            'description': s.description,
            'preconditions': s.preconditions,
            'steps': s.get_steps(),
            'expected_result': s.expected_result,
            'type': s.scenario_type,
            'priority': s.priority,
            'user_rating': s.user_rating,
        }
        for s in session.test_scenarios.all()
    ]


def _serialize_conditions(session):
    return [
        {
            'db_id': c.id,
            'condition_id': c.condition_id,
            'description': c.description,
            'type': c.condition_type,
            'priority': c.priority,
        }
        for c in session.test_conditions.all()
    ]


def _serialize_gaps(session):
    return [
        {
            'db_id': g.id,
            'issue_id': g.issue_id,
            'issue_type': g.issue_type,
            'description': g.description,
            'suggested_clarification': g.suggested_clarification,
        }
        for g in session.gaps.all()
    ]


def _serialize_quality_assessment(session):
    return {
        'clarity_score': session.clarity_score,
        'completeness_score': session.completeness_score,
        'testability_score': session.testability_score,
        'overall_score': session.accuracy_score,
        'severity': session.severity or 'Medium',
        'positive_aspects': [
            {'db_id': f.id, 'text': f.message}
            for f in session.feedback_items.filter(feedback_type='positive')
        ],
        'warnings': [
            {'db_id': f.id, 'text': f.message}
            for f in session.feedback_items.filter(feedback_type='warning')
        ],
    }


def _build_session_response(session, requirement_info=None, requires_decision=False, suggested=''):
    """Build the standard JSON response payload for a session."""
    info = requirement_info or session.get_extracted_info() or {'requirement_id': session.requirement_id or 'REQ-001'}
    return {
        'session_id': session.id,
        'requirements_text': session.requirements_text,
        'score': session.accuracy_score,
        'score_label': session.get_score_label(),
        'score_color': session.get_score_color(),
        'requirement_info': info,
        'quality_assessment': _serialize_quality_assessment(session),
        'test_conditions': _serialize_conditions(session),
        'gaps': _serialize_gaps(session),
        'scenarios': _serialize_scenarios(session),
        'requires_user_decision': requires_decision,
        'suggested_requirement': suggested,
        'team_notes': session.team_notes,
    }


def _get_accessible_session(session_id, user):
    """Return session if user is the owner or a workspace member, else raise Http404."""
    session = get_object_or_404(AnalysisSession, id=session_id)
    if session.user == user:
        return session
    if session.workspace and WorkspaceMembership.objects.filter(
        workspace=session.workspace, user=user
    ).exists():
        return session
    raise Http404()


# ─────────────────────────────────────────────
# Load session
# ─────────────────────────────────────────────

@login_required
def load_session_view(request, session_id):
    session = _get_accessible_session(session_id, request.user)
    return JsonResponse(_build_session_response(session))


# ─────────────────────────────────────────────
# Live workspace state (lightweight polling endpoint)
# ─────────────────────────────────────────────

@login_required
def workspace_state_view(request, workspace_id):
    """
    Cheap JSON snapshot of a shared workspace for near-real-time polling:
    the session list (for the sidebar) and the current member list.
    Makes no LLM calls. Only accessible to members of the workspace.
    """
    workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
    if not WorkspaceMembership.objects.filter(
        workspace=workspace, user=request.user
    ).exists():
        raise Http404()

    sessions = AnalysisSession.objects.filter(
        workspace=workspace
    ).order_by('-is_pinned', '-created_at')

    members = list(
        WorkspaceMembership.objects.filter(workspace=workspace)
        .select_related('user')
        .values_list('user__username', flat=True)
    )

    return JsonResponse({
        'sessions': [
            {
                'id': s.id,
                'title': s.title,
                'score': s.accuracy_score,
                'color': s.get_score_color(),
                'pinned': s.is_pinned,
            }
            for s in sessions
        ],
        'members': members,
        'member_count': len(members),
        'owner': workspace.owner.username,
    })


@login_required
def my_workspaces_state_view(request):
    """Cheap JSON snapshot of the user's team workspaces for live sidebar refresh."""
    return JsonResponse({'workspaces': [
        {'workspace_id': w.workspace_id, 'name': w.name, 'member_count': w.member_count}
        for w in _my_workspaces(request.user)
    ]})


# ─────────────────────────────────────────────
# Analyze requirements — Phase 3 QA redesign
# ─────────────────────────────────────────────

@login_required
@require_POST
def analyze_view(request):
    try:
        body = json.loads(request.body)
        requirements_text = body.get('requirements_text', '').strip()
        workspace_id = body.get('workspace_id', '').strip()
    except (json.JSONDecodeError, AttributeError):
        requirements_text = request.POST.get('requirements_text', '').strip()
        workspace_id = ''

    if not requirements_text:
        return JsonResponse({'error': 'Requirements text is required.'}, status=400)

    # Resolve workspace (if any)
    workspace = None
    if workspace_id:
        try:
            ws = Workspace.objects.get(workspace_id=workspace_id)
            if WorkspaceMembership.objects.filter(workspace=ws, user=request.user).exists():
                workspace = ws
        except Workspace.DoesNotExist:
            pass

    EDITED_PREFIX = 'EDITED:'
    if requirements_text.startswith(EDITED_PREFIX):
        edited_text = requirements_text[len(EDITED_PREFIX):].strip()
        try:
            qa_result = QAAnalyzer().analyze(edited_text)
        except GeminiQuotaExceeded as e:
            return JsonResponse({'error': str(e)}, status=429)
        session = TrainingDataManager().save_qa_session(
            user=request.user,
            requirements_text=edited_text,
            qa_result=qa_result,
        )
        if workspace:
            session.workspace = workspace
            session.save(update_fields=['workspace'])
        resp = _build_session_response(session, requirement_info=qa_result['requirement_info'])
        resp['system_action'] = 'analysis_and_generation'
        return JsonResponse(resp)

    try:
        qa_result = QAAnalyzer().analyze(requirements_text)
    except GeminiQuotaExceeded as e:
        return JsonResponse({'error': str(e)}, status=429)

    session = TrainingDataManager().save_qa_session(
        user=request.user,
        requirements_text=requirements_text,
        qa_result=qa_result,
    )
    if workspace:
        session.workspace = workspace
        session.save(update_fields=['workspace'])

    overall_score = session.accuracy_score
    suggested = qa_result.get('suggested_requirement', '')
    requires_improvement = overall_score < 80

    return JsonResponse(_build_session_response(
        session,
        requirement_info=qa_result['requirement_info'],
        requires_decision=requires_improvement,
        suggested=suggested if requires_improvement else '',
    ))


# ─────────────────────────────────────────────
# Section 6: Generate detailed test cases on demand
# ─────────────────────────────────────────────

@login_required
@require_POST
def generate_detailed_cases_view(request):
    try:
        body = json.loads(request.body)
        session_id = body.get('session_id')
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not session_id:
        return JsonResponse({'error': 'session_id is required.'}, status=400)

    session = _get_accessible_session(session_id, request.user)

    scenarios = [
        {
            'scenario_id': s.scenario_id,
            'description': s.description,
            'preconditions': s.preconditions,
            'expected_result': s.expected_result,
        }
        for s in session.test_scenarios.all()
    ]

    if not scenarios:
        return JsonResponse({'error': 'No scenarios found for this session.'}, status=400)

    try:
        detailed_cases = DetailedCaseGenerator().generate(scenarios)
    except GeminiQuotaExceeded as e:
        return JsonResponse({'error': str(e)}, status=429)
    TrainingDataManager().save_detailed_cases(session_id, detailed_cases)

    return JsonResponse({'session_id': session_id, 'detailed_cases': detailed_cases})


# ─────────────────────────────────────────────
# Update a single output field (feedback / gap / condition / scenario)
# ─────────────────────────────────────────────

@login_required
@require_POST
def update_output_field_view(request):
    try:
        body = json.loads(request.body)
        model_name = body.get('model')
        db_id = body.get('db_id')
        field = body.get('field')
        value = body.get('value', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request.'}, status=400)

    if not all([model_name, db_id, field]):
        return JsonResponse({'error': 'model, db_id, field are required.'}, status=400)

    def _has_access(session):
        if session.user == request.user:
            return True
        return session.workspace and WorkspaceMembership.objects.filter(
            workspace=session.workspace, user=request.user
        ).exists()

    if model_name == 'scenario':
        obj = get_object_or_404(TestScenario, id=db_id)
        if not _has_access(obj.session):
            return JsonResponse({'error': 'Access denied.'}, status=403)
        allowed = ('description', 'preconditions', 'expected_result')
        if field not in allowed:
            return JsonResponse({'error': 'Invalid field.'}, status=400)
        setattr(obj, field, value)
        obj.save(update_fields=[field])

    elif model_name == 'feedback':
        obj = get_object_or_404(FeedbackItem, id=db_id)
        if not _has_access(obj.session):
            return JsonResponse({'error': 'Access denied.'}, status=403)
        if field != 'message':
            return JsonResponse({'error': 'Invalid field.'}, status=400)
        obj.message = value
        obj.save(update_fields=['message'])

    elif model_name == 'gap':
        obj = get_object_or_404(RequirementGap, id=db_id)
        if not _has_access(obj.session):
            return JsonResponse({'error': 'Access denied.'}, status=403)
        allowed = ('description', 'suggested_clarification')
        if field not in allowed:
            return JsonResponse({'error': 'Invalid field.'}, status=400)
        setattr(obj, field, value)
        obj.save(update_fields=[field])

    elif model_name == 'condition':
        obj = get_object_or_404(TestCondition, id=db_id)
        if not _has_access(obj.session):
            return JsonResponse({'error': 'Access denied.'}, status=403)
        if field != 'description':
            return JsonResponse({'error': 'Invalid field.'}, status=400)
        obj.description = value
        obj.save(update_fields=['description'])

    else:
        return JsonResponse({'error': 'Invalid model.'}, status=400)

    return JsonResponse({'status': 'ok'})


# ─────────────────────────────────────────────
# Update team notes on a session
# ─────────────────────────────────────────────

@login_required
@require_POST
def update_team_notes_view(request):
    try:
        body = json.loads(request.body)
        session_id = body.get('session_id')
        notes = body.get('notes', '')
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request.'}, status=400)

    if not session_id:
        return JsonResponse({'error': 'session_id is required.'}, status=400)

    session = _get_accessible_session(session_id, request.user)
    session.team_notes = notes
    session.save(update_fields=['team_notes'])
    return JsonResponse({'status': 'ok'})


# ─────────────────────────────────────────────
# Delete all history
# ─────────────────────────────────────────────

@login_required
@require_POST
def delete_history_view(request):
    try:
        body = json.loads(request.body)
        workspace_id = body.get('workspace_id', '').strip()
    except (json.JSONDecodeError, AttributeError):
        workspace_id = ''

    if workspace_id:
        workspace = get_object_or_404(Workspace, workspace_id=workspace_id)
        if workspace.owner != request.user:
            return JsonResponse({'error': 'Only the workspace owner can clear history.'}, status=403)
        AnalysisSession.objects.filter(workspace=workspace).delete()
    else:
        AnalysisSession.objects.filter(user=request.user, workspace__isnull=True).delete()

    return JsonResponse({'status': 'cleared', 'message': 'History cleared successfully.'})


# ─────────────────────────────────────────────
# Toggle pin
# ─────────────────────────────────────────────

@login_required
@require_POST
def toggle_pin_view(request):
    try:
        body = json.loads(request.body)
        session_id = body.get('session_id')
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not session_id:
        return JsonResponse({'error': 'session_id is required.'}, status=400)

    session = _get_accessible_session(session_id, request.user)
    session.is_pinned = not session.is_pinned
    session.save()

    return JsonResponse({
        'system_action': 'chat_pinned',
        'session_id': session.id,
        'chat_metadata': {'is_pinned': session.is_pinned},
    })


# ─────────────────────────────────────────────
# Rename chat
# ─────────────────────────────────────────────

@login_required
@require_POST
def rename_chat_view(request):
    try:
        body = json.loads(request.body)
        session_id = body.get('session_id')
        new_title = body.get('new_title', '').strip()
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not session_id or not new_title:
        return JsonResponse({'error': 'session_id and new_title are required.'}, status=400)

    session = _get_accessible_session(session_id, request.user)
    session.title = new_title[:200]
    session.save()

    return JsonResponse({'system_action': 'chat_renamed', 'chat_title': session.title, 'session_id': session.id})


# ─────────────────────────────────────────────
# Delete current chat
# ─────────────────────────────────────────────

@login_required
@require_POST
def delete_current_chat_view(request):
    try:
        body = json.loads(request.body)
        session_id = body.get('session_id')
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request body.'}, status=400)

    if not session_id:
        return JsonResponse({'error': 'session_id is required.'}, status=400)

    session = _get_accessible_session(session_id, request.user)

    # Team (workspace) chats can only be deleted by the workspace owner.
    if session.workspace and session.workspace.owner != request.user:
        return JsonResponse(
            {'error': 'Only the workspace owner can delete team chats.'}, status=403
        )

    session.delete()

    return JsonResponse({'system_action': 'chat_deleted', 'session_id': session_id})


# ─────────────────────────────────────────────
# Rate a scenario
# ─────────────────────────────────────────────

@login_required
@require_POST
def rate_scenario_view(request):
    try:
        body = json.loads(request.body)
        db_id  = body.get('db_id')
        rating = body.get('rating', '')
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request.'}, status=400)

    if rating not in ('useful', 'not_useful', ''):
        return JsonResponse({'error': 'Invalid rating value.'}, status=400)

    scenario = get_object_or_404(TestScenario, id=db_id)
    session  = scenario.session

    if session.user != request.user:
        if not (session.workspace and WorkspaceMembership.objects.filter(
            workspace=session.workspace, user=request.user
        ).exists()):
            return JsonResponse({'error': 'Access denied.'}, status=403)

    scenario.user_rating = '' if scenario.user_rating == rating else rating
    scenario.save()

    return JsonResponse({'status': 'ok', 'db_id': scenario.id, 'rating': scenario.user_rating})


# ─────────────────────────────────────────────
# Re-analyze an existing session
# ─────────────────────────────────────────────

@login_required
@require_POST
def reanalyze_view(request):
    try:
        body = json.loads(request.body)
        session_id = body.get('session_id')
    except (json.JSONDecodeError, AttributeError):
        return JsonResponse({'error': 'Invalid request.'}, status=400)

    session = _get_accessible_session(session_id, request.user)

    try:
        qa_result = QAAnalyzer().analyze(session.requirements_text)
    except GeminiQuotaExceeded as e:
        return JsonResponse({'error': str(e)}, status=429)

    TrainingDataManager().reanalyze_session(session, qa_result)

    return JsonResponse(_build_session_response(
        session,
        requirement_info=qa_result['requirement_info'],
        requires_decision=session.accuracy_score < 80,
        suggested=session.suggested_requirement,
    ))


# ─────────────────────────────────────────────
# Export helpers — collect the full QA report once for all formats
# ─────────────────────────────────────────────

_ANALYSIS_FIELDS = [
    ('Actors', 'actors'),
    ('Actions', 'actions'),
    ('Business Rules', 'business_rules'),
    ('Constraints', 'constraints'),
    ('Validation Rules', 'validation_rules'),
    ('Error Handling', 'error_handling'),
    ('Non-Functional', 'non_functional'),
]


def _scenario_detailed_case(scenario):
    try:
        return scenario.detailed_case
    except DetailedTestCase.DoesNotExist:
        return None


def _collect_export_data(session):
    info = session.get_extracted_info() or {}
    scenarios = list(session.test_scenarios.all())
    detailed = [(s.scenario_id, dc) for s in scenarios
                if (dc := _scenario_detailed_case(s)) is not None]
    return {
        'title': session.title,
        'requirement_id': session.requirement_id or info.get('requirement_id', 'REQ-001'),
        'requirements_text': session.requirements_text,
        'generated_by': session.user.username,
        'created_at': session.created_at,
        'clarity': session.clarity_score,
        'completeness': session.completeness_score,
        'testability': session.testability_score,
        'overall': session.accuracy_score,
        'severity': session.severity or 'Medium',
        'label': session.get_score_label(),
        'analysis': info,
        'strengths': [f.message for f in session.feedback_items.filter(feedback_type='positive')],
        'warnings': [f.message for f in session.feedback_items.filter(feedback_type='warning')],
        'conditions': list(session.test_conditions.all()),
        'gaps': list(session.gaps.all()),
        'scenarios': scenarios,
        'detailed': detailed,
        'team_notes': session.team_notes,
    }


# ─────────────────────────────────────────────
# Export — CSV (full report, multi-section)
# ─────────────────────────────────────────────

@login_required
def export_csv_view(request, session_id):
    import csv
    session = _get_accessible_session(session_id, request.user)
    d = _collect_export_data(session)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="qa_report_{session_id}.csv"'
    w = csv.writer(response)
    blank = lambda: w.writerow([])

    # 1. Summary
    w.writerow(['SECTION', 'Summary'])
    w.writerow(['Project', d['title']])
    w.writerow(['Requirement ID', d['requirement_id']])
    w.writerow(['Overall Score', f"{d['overall']}% ({d['label']})"])
    w.writerow(['Clarity', f"{d['clarity']}%"])
    w.writerow(['Completeness', f"{d['completeness']}%"])
    w.writerow(['Testability', f"{d['testability']}%"])
    w.writerow(['Severity', d['severity']])
    w.writerow(['Generated By', d['generated_by']])
    w.writerow(['Date', d['created_at'].strftime('%Y-%m-%d %H:%M')])
    blank()

    # 2. Requirement text
    w.writerow(['SECTION', 'Requirement'])
    w.writerow([d['requirements_text']])
    blank()

    # 3. Requirement Analysis
    w.writerow(['SECTION', 'Requirement Analysis'])
    for label, key in _ANALYSIS_FIELDS:
        w.writerow([label, '; '.join(str(x) for x in d['analysis'].get(key, []))])
    blank()

    # 4. Quality findings
    w.writerow(['SECTION', 'Strengths'])
    for x in d['strengths']:
        w.writerow([x])
    blank()
    w.writerow(['SECTION', 'Warnings'])
    for x in d['warnings']:
        w.writerow([x])
    blank()

    # 5. Test Conditions
    w.writerow(['SECTION', 'Test Conditions'])
    w.writerow(['Condition ID', 'Description', 'Type', 'Priority'])
    for c in d['conditions']:
        w.writerow([c.condition_id, c.description, c.condition_type, c.priority])
    blank()

    # 6. Review Findings
    w.writerow(['SECTION', 'Review Findings'])
    w.writerow(['Issue ID', 'Issue Type', 'Description', 'Suggested Clarification'])
    for g in d['gaps']:
        w.writerow([g.issue_id, g.issue_type, g.description, g.suggested_clarification])
    blank()

    # 7. Test Scenarios
    w.writerow(['SECTION', 'Test Scenarios'])
    w.writerow(['ID', 'Req Ref', 'Cond Ref', 'Description', 'Preconditions',
                'Steps', 'Expected Result', 'Type', 'Priority', 'Rating'])
    for s in d['scenarios']:
        w.writerow([s.scenario_id, s.requirement_ref, s.condition_ref, s.description,
                    s.preconditions, ' | '.join(s.get_steps()), s.expected_result,
                    s.scenario_type, s.priority, s.user_rating or '-'])
    blank()

    # 8. Detailed Test Cases (only if generated)
    if d['detailed']:
        w.writerow(['SECTION', 'Detailed Test Cases'])
        w.writerow(['Scenario ID', 'Test Data', 'Steps', 'Expected Results', 'Postconditions'])
        for sid, dc in d['detailed']:
            w.writerow([sid, dc.test_data, ' | '.join(dc.get_steps()),
                        dc.expected_results, dc.postconditions])
        blank()

    # 9. Team Notes (only if any)
    if d['team_notes']:
        w.writerow(['SECTION', 'Team Notes'])
        w.writerow([d['team_notes']])

    return response


# ─────────────────────────────────────────────
# Export — PDF
# ─────────────────────────────────────────────

@login_required
def export_pdf_view(request, session_id):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    session = _get_accessible_session(session_id, request.user)
    d = _collect_export_data(session)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=2*cm, bottomMargin=2*cm
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('ReportTitle', parent=styles['Title'], fontSize=16, spaceAfter=6)
    heading_style = ParagraphStyle('ReportHeading', parent=styles['Heading2'], fontSize=12, spaceBefore=8, spaceAfter=4)
    normal_style = ParagraphStyle('ReportNormal', parent=styles['Normal'], fontSize=9, leading=12)
    header_cell_style = ParagraphStyle('HeaderCell', parent=styles['Normal'], fontSize=8, leading=10,
                                       textColor=colors.white, fontName='Helvetica-Bold')
    body_cell_style = ParagraphStyle('BodyCell', parent=styles['Normal'], fontSize=7.5, leading=10,
                                      textColor=colors.black, wordWrap='CJK')

    table_style = TableStyle([
        ('BACKGROUND',    (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR',     (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('GRID',          (0, 0), (-1, -1), 0.4, colors.HexColor('#cccccc')),
        ('VALIGN',        (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING',    (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING',   (0, 0), (-1, -1), 4),
        ('RIGHTPADDING',  (0, 0), (-1, -1), 4),
    ])

    def esc(t):
        return (str(t).replace('&', '&amp;').replace('<', '&lt;')
                .replace('>', '&gt;').replace('\n', '<br/>'))

    def make_table(headers, rows, col_widths):
        data = [[Paragraph(h, header_cell_style) for h in headers]]
        for row in rows:
            data.append([c if isinstance(c, Paragraph) else Paragraph(esc(c), body_cell_style) for c in row])
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(table_style)
        return t

    story = []

    # ── Header / summary ──
    story.append(Paragraph('QA Analysis &amp; Test Scenario Report', title_style))
    story.append(Paragraph(f'Project: {esc(d["title"])}', heading_style))
    story.append(Paragraph(
        f'Requirement ID: {esc(d["requirement_id"])} &nbsp;|&nbsp; Generated by: {esc(d["generated_by"])} '
        f'&nbsp;|&nbsp; {d["created_at"].strftime("%d %b %Y %H:%M")}', normal_style))
    story.append(Spacer(1, 0.3*cm))
    story.append(make_table(
        ['Clarity', 'Completeness', 'Testability', 'Overall', 'Severity'],
        [[f'{d["clarity"]}%', f'{d["completeness"]}%', f'{d["testability"]}%',
          f'{d["overall"]}% ({d["label"]})', d['severity']]],
        [3.4*cm]*5))
    story.append(Spacer(1, 0.4*cm))

    # ── 1. Requirement ──
    story.append(Paragraph('1. Requirement', heading_style))
    story.append(Paragraph(esc(d['requirements_text']), normal_style))

    # ── 2. Requirement Analysis ──
    analysis_rows = [[label, '; '.join(str(x) for x in d['analysis'].get(key, []))]
                     for label, key in _ANALYSIS_FIELDS if d['analysis'].get(key)]
    if analysis_rows:
        story.append(Paragraph('2. Requirement Analysis', heading_style))
        story.append(make_table(['Field', 'Details'], analysis_rows, [4*cm, 13*cm]))

    # ── 3. Quality Assessment ──
    if d['strengths'] or d['warnings']:
        story.append(Paragraph('3. Quality Assessment', heading_style))
        for s in d['strengths']:
            story.append(Paragraph(f'&#10003; {esc(s)}', normal_style))
        for wn in d['warnings']:
            story.append(Paragraph(f'&#9650; {esc(wn)}', normal_style))

    # ── 4. Test Conditions ──
    if d['conditions']:
        story.append(Paragraph('4. Test Conditions', heading_style))
        rows = [[c.condition_id, c.description, c.condition_type, c.priority] for c in d['conditions']]
        story.append(make_table(['ID', 'Condition', 'Type', 'Priority'], rows,
                                [1.6*cm, 10.4*cm, 2.6*cm, 2.4*cm]))

    # ── 5. Review Findings ──
    if d['gaps']:
        story.append(Paragraph('5. Requirement Review Findings', heading_style))
        rows = [[g.issue_id, g.issue_type, g.description, g.suggested_clarification] for g in d['gaps']]
        story.append(make_table(['ID', 'Type', 'Description', 'Suggested Clarification'], rows,
                                [1.4*cm, 2.8*cm, 6.4*cm, 6.4*cm]))

    # ── 6. Test Scenarios ──
    if d['scenarios']:
        story.append(Paragraph('6. Traceable Test Scenarios', heading_style))
        rows = []
        for s in d['scenarios']:
            steps_html = '<br/>'.join(f'{i+1}. {esc(step)}' for i, step in enumerate(s.get_steps()))
            meta = (f'<font size=6 color="#888888">Req: {esc(s.requirement_ref or "-")} &middot; '
                    f'Cond: {esc(s.condition_ref or "-")} &middot; {esc(s.scenario_type)} &middot; '
                    f'Rating: {esc(s.user_rating or "-")}</font><br/>')
            desc = Paragraph(meta + esc(s.description), body_cell_style)
            rows.append([s.scenario_id, desc, s.preconditions, steps_html, s.expected_result, s.priority])
        story.append(make_table(
            ['ID', 'Description', 'Pre-Conditions', 'Test Steps', 'Expected Result', 'Priority'],
            rows, [1.3*cm, 4.2*cm, 3.0*cm, 3.9*cm, 3.1*cm, 1.5*cm]))

    # ── 7. Detailed Test Cases (only if generated) ──
    if d['detailed']:
        story.append(Paragraph('7. Detailed Test Cases', heading_style))
        rows = []
        for sid, dc in d['detailed']:
            steps_html = '<br/>'.join(f'{i+1}. {esc(step)}' for i, step in enumerate(dc.get_steps()))
            rows.append([sid, dc.test_data, steps_html, dc.expected_results, dc.postconditions])
        story.append(make_table(
            ['Scenario', 'Test Data', 'Steps', 'Expected Results', 'Postconditions'],
            rows, [1.8*cm, 3.4*cm, 4.4*cm, 3.7*cm, 3.7*cm]))

    # ── 8. Team Notes (only if any) ──
    if d['team_notes']:
        story.append(Paragraph('8. Team Notes', heading_style))
        story.append(Paragraph(esc(d['team_notes']), normal_style))

    doc.build(story)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="qa_report_{session_id}.pdf"'
    return response


# ─────────────────────────────────────────────
# Export — Excel
# ─────────────────────────────────────────────

@login_required
def export_excel_view(request, session_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    session = _get_accessible_session(session_id, request.user)
    d = _collect_export_data(session)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2c3e50', end_color='2c3e50', fill_type='solid')
    label_font = Font(bold=True)
    top_wrap = Alignment(vertical='top', wrap_text=True)
    center_wrap = Alignment(horizontal='center', vertical='top', wrap_text=True)

    def header_row(ws, headers, row=1):
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_wrap

    def set_widths(ws, widths):
        for i, wdt in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wdt

    def wrap_sheet(ws, start_row=2):
        for row in ws.iter_rows(min_row=start_row):
            for cell in row:
                cell.alignment = top_wrap

    wb = Workbook()

    # ── Sheet 1: Summary ──
    s1 = wb.active
    s1.title = 'Summary'
    summary = [
        ('Project', d['title']),
        ('Requirement ID', d['requirement_id']),
        ('Overall Score', f"{d['overall']}% ({d['label']})"),
        ('Clarity', f"{d['clarity']}%"),
        ('Completeness', f"{d['completeness']}%"),
        ('Testability', f"{d['testability']}%"),
        ('Severity', d['severity']),
        ('Generated By', d['generated_by']),
        ('Date', d['created_at'].strftime('%Y-%m-%d %H:%M')),
        ('Requirement', d['requirements_text']),
        ('Strengths', '\n'.join(d['strengths'])),
        ('Warnings', '\n'.join(d['warnings'])),
    ]
    for r, (label, value) in enumerate(summary, start=1):
        c1 = s1.cell(row=r, column=1, value=label); c1.font = label_font; c1.alignment = top_wrap
        s1.cell(row=r, column=2, value=value).alignment = top_wrap
    set_widths(s1, [18, 80])

    # ── Sheet 2: Requirement Analysis ──
    s2 = wb.create_sheet('Requirement Analysis')
    header_row(s2, ['Field', 'Details'])
    for r, (label, key) in enumerate(_ANALYSIS_FIELDS, start=2):
        s2.cell(row=r, column=1, value=label)
        s2.cell(row=r, column=2, value='\n'.join(str(x) for x in d['analysis'].get(key, [])))
    wrap_sheet(s2)
    set_widths(s2, [22, 90])

    # ── Sheet 3: Test Conditions ──
    s3 = wb.create_sheet('Test Conditions')
    header_row(s3, ['Condition ID', 'Description', 'Type', 'Priority'])
    for r, c in enumerate(d['conditions'], start=2):
        s3.cell(row=r, column=1, value=c.condition_id)
        s3.cell(row=r, column=2, value=c.description)
        s3.cell(row=r, column=3, value=c.condition_type)
        s3.cell(row=r, column=4, value=c.priority)
    wrap_sheet(s3)
    set_widths(s3, [14, 70, 14, 12])

    # ── Sheet 4: Review Findings ──
    s4 = wb.create_sheet('Review Findings')
    header_row(s4, ['Issue ID', 'Issue Type', 'Description', 'Suggested Clarification'])
    for r, g in enumerate(d['gaps'], start=2):
        s4.cell(row=r, column=1, value=g.issue_id)
        s4.cell(row=r, column=2, value=g.issue_type)
        s4.cell(row=r, column=3, value=g.description)
        s4.cell(row=r, column=4, value=g.suggested_clarification)
    wrap_sheet(s4)
    set_widths(s4, [12, 18, 45, 45])

    # ── Sheet 5: Test Scenarios ──
    s5 = wb.create_sheet('Test Scenarios')
    header_row(s5, ['ID', 'Req Ref', 'Cond Ref', 'Description', 'Preconditions',
                    'Test Steps', 'Expected Result', 'Type', 'Priority', 'Rating'])
    for r, s in enumerate(d['scenarios'], start=2):
        values = [s.scenario_id, s.requirement_ref, s.condition_ref, s.description,
                  s.preconditions, '\n'.join(s.get_steps()), s.expected_result,
                  s.scenario_type, s.priority, s.user_rating or '-']
        for col, value in enumerate(values, start=1):
            s5.cell(row=r, column=col, value=value)
    wrap_sheet(s5)
    set_widths(s5, [10, 12, 10, 32, 28, 40, 32, 10, 10, 10])

    # ── Sheet 6: Detailed Cases (only if generated) ──
    if d['detailed']:
        s6 = wb.create_sheet('Detailed Cases')
        header_row(s6, ['Scenario ID', 'Test Data', 'Steps', 'Expected Results', 'Postconditions'])
        for r, (sid, dc) in enumerate(d['detailed'], start=2):
            s6.cell(row=r, column=1, value=sid)
            s6.cell(row=r, column=2, value=dc.test_data)
            s6.cell(row=r, column=3, value='\n'.join(dc.get_steps()))
            s6.cell(row=r, column=4, value=dc.expected_results)
            s6.cell(row=r, column=5, value=dc.postconditions)
        wrap_sheet(s6)
        set_widths(s6, [14, 35, 45, 35, 30])

    # ── Sheet 7: Team Notes (only if any) ──
    if d['team_notes']:
        s7 = wb.create_sheet('Team Notes')
        s7.cell(row=1, column=1, value=d['team_notes']).alignment = top_wrap
        set_widths(s7, [100])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="qa_report_{session_id}.xlsx"'
    return response
